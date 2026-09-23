"""Arquivo de verificação (Excel): quantitativos por escola e conferência de cada aluno/nome da base.

É gerado automaticamente ao final de `gerar` e pode ser refeito a qualquer momento
(por exemplo, depois de gerar escola por escola):

    python -m capas_prova.verificar --planilha Base.xlsx --saida saida

Abas:
  Resumo por escola  alunos na base (total e por série), caderno 1 x 2, capas esperadas x geradas
  Sorteio por turma  quantos alunos de cada turma ficaram com o caderno 1 e com o 2
  Totais por caderno alunos e capas geradas por caderno (C0201, C0202, ...)
  Conferir           só os alunos que pedem atenção (capa faltando, RA repetido, nome suspeito)
  Alunos             todos os alunos da base: nome original x nome na capa, alterações, capas e códigos
  Legenda            explicação das colunas e dos status
"""
from __future__ import annotations

import argparse
import csv
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
import yaml

from .nomes import PARTICULAS_PADRAO, padronizar_nome

OK, CONFERIR, FALTA = "OK", "CONFERIR", "FALTA CAPA"


def _ler_manifesto(saida: Path) -> dict[tuple[str, str], dict]:
    arq = saida / "manifesto.csv"
    if not arq.exists():
        return {}
    with open(arq, newline="", encoding="utf-8-sig") as f:
        return {(r["ra"], r["caderno"]): r for r in csv.DictReader(f)}


def _serie_num(serie: str) -> str:
    m = re.search(r"\d+", serie)
    return f"{m.group()}º ano" if m else "sem série"


def montar_tabelas(cfg: dict, planilha: Path, saida: Path, ajustes: dict | None = None,
                   validacao: bool = False) -> dict[str, pd.DataFrame]:
    from .gerar import cadernos_do_aluno, ler_planilha  # import tardio (evita ciclo)

    from .sorteio import cadernos_da_serie, sortear

    ajustes = ajustes or {}
    df = ler_planilha(planilha, cfg)
    if cfg.get("sorteio", {}).get("ativo", True):
        sorteados, _ = sortear(df, cfg, saida, salvar=False)   # lê o sorteio salvo (não sorteia de novo)
        df["caderno"] = df["ra"].map(sorteados).fillna("")
    cfg_nomes = cfg.get("nomes", {})
    manifesto = _ler_manifesto(saida)

    linhas_por_ra = defaultdict(list)
    for idx, ra in df["ra"].items():
        linhas_por_ra[ra].append(idx + 2)

    alunos = []
    for idx, r in df.iterrows():
        n = idx + 2
        original = r["estudante"]
        if cfg_nomes.get("padronizar", True):
            nome, alteracoes, conferir = padronizar_nome(
                original, cfg_nomes.get("particulas_minusculas", PARTICULAS_PADRAO),
                cfg_nomes.get("remover_parenteses", True))
        else:
            nome, alteracoes, conferir = original, [], []
        so_letras = lambda s: re.sub(r"[\W_\d]", "", s)  # noqa: E731
        if so_letras(nome) != so_letras(original) and so_letras(nome).lower() == so_letras(original).lower():
            alteracoes.insert(0, "maiúsculas/minúsculas padronizadas")

        faltando_campos = [c for c in ("escola", "serie", "turma", "estudante", "ra") if not r.get(c, "")]
        if faltando_campos:
            conferir.append(f"campos vazios na base: {', '.join(faltando_campos)}")
        repetido = len(linhas_por_ra[r["ra"]]) > 1 and r["ra"]
        dono = not repetido or linhas_por_ra[r["ra"]][0] == n
        if repetido:
            outras = [str(x) for x in linhas_por_ra[r["ra"]] if x != n]
            conferir.append(f"RA repetido na base (também na linha {', '.join(outras)})")

        esperados = cadernos_do_aluno(r.to_dict(), cfg) if not faltando_campos else []
        if not esperados and not faltando_campos:
            conferir.append(f"nenhum caderno configurado para a série '{r['serie']}'")
        geradas, faltam, codigos, ajustes_txt = [], [], [], []
        for cad in esperados:
            chave = (r["ra"], cad)
            m = manifesto.get(chave)
            existe = bool(dono and m and (saida / m["arquivo"]).exists())
            if existe:
                geradas.append(cad)
                codigos.append(f"{cad}: {m['codigo_cartao']}")
            elif dono:
                faltam.append(cad)
            aj = ajustes.get(chave) or (m or {}).get("ajustes_texto", "") if dono else ""
            if aj:
                ajustes_txt.append(f"{cad}: {aj}")

        if faltam and not validacao:
            status = FALTA
        elif conferir or not dono:
            status = CONFERIR
        else:
            status = OK
        alunos.append({
            "Linha na base": n,
            "Escola": r["escola"],
            "Série": r["serie"],
            "Turma": r["turma"],
            "RA": r["ra"],
            "Nome na base": original,
            "Nome na capa": nome,
            "Alterações no nome": "; ".join(alteracoes),
            "Conferir": "; ".join(conferir),
            "Caderno sorteado": ", ".join(esperados) if dono else "(capa gerada pela 1ª linha do RA)",
            "Capas geradas": ", ".join(geradas),
            "Capas faltando": ", ".join(faltam),
            "Códigos do cartão": "; ".join(codigos),
            "Ajustes de texto na capa": " | ".join(ajustes_txt),
            "Status": status,
            "_esperadas": len(esperados) if dono else 0,
            "_geradas": len(geradas),
            "_serie": _serie_num(r["serie"]),
            "_ra_unico": dono,
            "_num_caderno": (cadernos_da_serie(r["serie"], cfg).index(esperados[0]) + 1
                             if dono and esperados and esperados[0] in cadernos_da_serie(r["serie"], cfg) else 0),
        })

    tab = pd.DataFrame(alunos)
    series = sorted(tab["_serie"].unique(), key=lambda s: (len(s), s))
    resumo = []
    for escola, g in tab.groupby("Escola", sort=True):
        linha = {"Escola": escola, "Alunos na base (linhas)": len(g)}
        for s in series:
            linha[f"Alunos {s}"] = int((g["_serie"] == s).sum())
        c1, c2 = int((g["_num_caderno"] == 1).sum()), int((g["_num_caderno"] == 2).sum())
        linha |= {"Caderno 1 (x01)": c1, "Caderno 2 (x02)": c2, "Diferença cadernos": abs(c1 - c2)}
        esp, ger = int(g["_esperadas"].sum()), int(g["_geradas"].sum())
        linha |= {
            "Capas esperadas": esp,
            "Capas geradas": ger,
            "Capas faltando": esp - ger,
            "% geradas": round(100 * ger / esp, 1) if esp else 0.0,
            "Nomes alterados": int((g["Alterações no nome"] != "").sum()),
            "Alunos a conferir": int((g["Status"] == CONFERIR).sum()),
            "RA repetido": int(g["Conferir"].str.contains("RA repetido").sum()),
            "Textos reduzidos/quebrados": int((g["Ajustes de texto na capa"] != "").sum()),
            "Status": ("OK" if esp == ger else ("PENDENTE" if ger == 0 else "INCOMPLETO")) if not validacao else "PRÉVIA",
        }
        resumo.append(linha)
    resumo = pd.DataFrame(resumo)
    total = {c: (resumo[c].sum() if resumo[c].dtype != object else "") for c in resumo.columns}
    total["Escola"] = "TOTAL"
    total["% geradas"] = round(100 * total["Capas geradas"] / total["Capas esperadas"], 1) if total["Capas esperadas"] else 0.0
    total["Status"] = ""
    total["Diferença cadernos"] = abs(total["Caderno 1 (x01)"] - total["Caderno 2 (x02)"])
    resumo = pd.concat([resumo, pd.DataFrame([total])], ignore_index=True)

    # equilíbrio do sorteio por turma e por caderno
    por_turma = []
    for (esc, ser, tur), g in tab[tab["_num_caderno"] > 0].groupby(["Escola", "Série", "Turma"], sort=True):
        c1, c2 = int((g["_num_caderno"] == 1).sum()), int((g["_num_caderno"] == 2).sum())
        por_turma.append({"Escola": esc, "Série": ser, "Turma": tur, "Alunos": len(g),
                          "Caderno 1 (x01)": c1, "Caderno 2 (x02)": c2, "Diferença": abs(c1 - c2)})
    por_turma = pd.DataFrame(por_turma)
    por_caderno = (tab[tab["Caderno sorteado"].str.match(r"^C\d{4}$")]
                   .groupby("Caderno sorteado").agg(Alunos=("RA", "count"), **{"Capas geradas": ("_geradas", "sum")})
                   .reset_index().rename(columns={"Caderno sorteado": "Caderno"}))

    visiveis = [c for c in tab.columns if not c.startswith("_")]
    conferir = tab[tab["Status"] != OK][visiveis]
    legenda = pd.DataFrame([
        ("Status OK", "todas as capas esperadas foram geradas e o nome não tem pendências"),
        ("Status CONFERIR", "capa gerada, mas algo precisa ser olhado por uma pessoa (ver coluna Conferir)"),
        ("Status FALTA CAPA", "o aluno está na base, mas falta capa de algum caderno (ainda não gerada ou erro)"),
        ("Resumo: PENDENTE / INCOMPLETO", "escola sem nenhuma capa gerada / com parte das capas"),
        ("Resumo: PRÉVIA", "arquivo feito com --validar: nenhuma capa foi gerada, só a conferência da base"),
        ("Nome na capa", "nome padronizado: inicial maiúscula, partículas (de, da, do, das, dos, e) minúsculas"),
        ("Alterações no nome", "o que a padronização mudou automaticamente (não exige ação)"),
        ("Conferir", "o que a padronização NÃO resolve sozinha: letra isolada/abreviação, texto entre "
                     "parênteses, caracteres incomuns, RA repetido etc."),
        ("Ajustes de texto na capa", "campos que ficaram com letra menor ou quebrados em 2 linhas para caber"),
        ("Capas geradas", "conferido no manifesto E no disco (o PDF existe na pasta de saída)"),
        ("Caderno sorteado", "caderno do aluno, sorteado de forma equilibrada dentro da turma "
                             "(arquivo sorteio_cadernos.csv na pasta de saída)"),
        ("Diferença cadernos / Diferença", "|caderno 1 − caderno 2|: no máximo 1 por turma; por escola e no total "
                                           "fica 0 ou 1 por série"),
    ], columns=["Item", "Significado"])
    return {"Resumo por escola": resumo, "Sorteio por turma": por_turma, "Totais por caderno": por_caderno,
            "Conferir": conferir, "Alunos": tab[visiveis], "Legenda": legenda}


def _formatar(ws, largura_max=60):
    from openpyxl.styles import Alignment, Font, PatternFill

    cores = {OK: "C6EFCE", CONFERIR: "FFEB9C", FALTA: "FFC7CE", "INCOMPLETO": "FFEB9C", "PENDENTE": "FFC7CE"}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    cab = [c.value for c in ws[1]]
    for i, col in enumerate(ws.columns):
        amostra = [len(str(c.value)) for c in list(col)[:500] if c.value is not None]
        ws.column_dimensions[col[0].column_letter].width = min(largura_max, max([10] + amostra) + 2)
        if cab[i] == "Status":
            for c in list(col)[1:]:
                if c.value in cores:
                    c.fill = PatternFill("solid", fgColor=cores[c.value])
    ultima = ws[ws.max_row]
    if ultima[0].value == "TOTAL":
        for c in ultima:
            c.font = Font(bold=True)


def gerar_verificacao(cfg: dict, planilha: Path, saida: Path, ajustes: dict | None = None,
                      validacao: bool = False) -> Path:
    tabelas = montar_tabelas(cfg, planilha, saida, ajustes, validacao)
    saida.mkdir(parents=True, exist_ok=True)
    arq = saida / ("verificacao_previa.xlsx" if validacao else "verificacao.xlsx")
    with pd.ExcelWriter(arq, engine="openpyxl") as w:
        for nome, t in tabelas.items():
            t.to_excel(w, sheet_name=nome, index=False)
            _formatar(w.sheets[nome], largura_max=90 if nome == "Legenda" else 60)
    return arq


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--planilha", required=True, type=Path)
    ap.add_argument("--saida", required=True, type=Path)
    ap.add_argument("--config", type=Path, default=Path(__file__).resolve().parent.parent / "config.yaml")
    args = ap.parse_args(argv)
    cfg = yaml.safe_load(args.config.read_text(encoding="utf-8"))
    arq = gerar_verificacao(cfg, args.planilha, args.saida)
    resumo = pd.read_excel(arq, sheet_name="Resumo por escola")
    tot = resumo.iloc[-1]
    st = Counter(pd.read_excel(arq, sheet_name="Alunos")["Status"])
    print(f"Verificação: {arq}")
    print(f"  capas esperadas {tot['Capas esperadas']}, geradas {tot['Capas geradas']}, faltando {tot['Capas faltando']}")
    print(f"  alunos: {dict(st)}")


if __name__ == "__main__":
    main()
